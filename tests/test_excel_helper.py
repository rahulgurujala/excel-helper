import os

import openpyxl
import pandas as pd
import pytest

from openpyxl.styles import Font

from excel_helper import ExcelHelper


@pytest.fixture
def excel_file():
    filename = "test_excel.xlsx"
    yield filename
    if os.path.exists(filename):
        os.remove(filename)


@pytest.fixture
def excel_helper(excel_file):
    helper = ExcelHelper(excel_file)
    helper.create_new_workbook()
    return helper


def test_create_and_save_workbook(excel_helper: ExcelHelper, excel_file):
    excel_helper.save_workbook()
    assert os.path.exists(excel_file)


def test_write_cell_requires_open_workbook(excel_file):
    helper = ExcelHelper(excel_file)
    with pytest.raises(ValueError, match="Workbook is not open"):
        helper.write_cell(1, 1, "Test")


def test_write_and_read_cell(excel_helper: ExcelHelper):
    excel_helper.write_cell(1, 1, "Test")
    assert excel_helper.read_cell(1, 1) == "Test"


def test_write_and_read_row(excel_helper: ExcelHelper):
    test_data = ["A", "B", "C"]
    excel_helper.write_row(1, test_data)
    assert excel_helper.read_row(1) == test_data


def test_write_and_read_column(excel_helper: ExcelHelper):
    test_data = ["X", "Y", "Z"]
    excel_helper.write_column(1, test_data)
    assert excel_helper.read_column(1) == test_data


def test_write_and_read_range(excel_helper: ExcelHelper):
    test_data = [["1", "2"], ["3", "4"]]
    excel_helper.write_range(1, 1, test_data)
    assert excel_helper.read_range(1, 1, 2, 2) == test_data


def test_set_and_get_formula(excel_helper: ExcelHelper):
    formula = "=SUM(A1:A5)"
    excel_helper.set_formula(1, 1, formula)
    assert excel_helper.get_formula(1, 1) == formula


def test_sum_range(excel_helper: ExcelHelper):
    test_data = [[1], [2], [3], [4], [5]]
    excel_helper.write_range(1, 1, test_data)
    excel_helper.sum_range(1, 1, 5, 1, 6, 1)
    assert excel_helper.get_formula(6, 1) == "=SUM(A1:A5)"


def test_average_range(excel_helper: ExcelHelper):
    test_data = [[1], [2], [3], [4], [5]]
    excel_helper.write_range(1, 1, test_data)
    excel_helper.average_range(1, 1, 5, 1, 6, 1)
    assert excel_helper.get_formula(6, 1) == "=AVERAGE(A1:A5)"


def test_count_range(excel_helper: ExcelHelper):
    test_data = [[1], [2], [3], [4], [5]]
    excel_helper.write_range(1, 1, test_data)
    excel_helper.count_range(1, 1, 5, 1, 6, 1)
    assert excel_helper.get_formula(6, 1) == "=COUNT(A1:A5)"


def test_if_formula(excel_helper: ExcelHelper):
    excel_helper.if_formula(1, 1, "True", "False", 2, 1)
    assert excel_helper.get_formula(2, 1) == '=IF(A1, "True", "False")'


def test_vlookup(excel_helper: ExcelHelper):
    test_data = [["A", 1], ["B", 2], ["C", 3]]
    excel_helper.write_range(1, 1, test_data)
    excel_helper.write_cell(5, 1, "B")
    excel_helper.vlookup(5, 1, 1, 1, 3, 2, 2, 5, 2)
    assert excel_helper.get_formula(5, 2) == "=VLOOKUP(A5, A1:B3, 2, FALSE)"


def test_select_sheet(excel_helper: ExcelHelper):
    excel_helper.workbook.create_sheet("TestSheet")
    excel_helper.select_sheet("TestSheet")
    assert excel_helper.active_sheet.title == "TestSheet"


def test_select_sheet_requires_open_workbook(excel_file):
    helper = ExcelHelper(excel_file)
    with pytest.raises(ValueError, match="Workbook is not open"):
        helper.select_sheet("Sheet1")


def test_auto_fit_columns(excel_helper: ExcelHelper):
    test_data = [["Short", "A very long column header"]]
    excel_helper.write_range(1, 1, test_data)
    excel_helper.auto_fit_columns()
    assert (
        excel_helper.active_sheet.column_dimensions["A"].width
        < excel_helper.active_sheet.column_dimensions["B"].width
    )


def test_auto_fit_columns_with_numeric_values(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [[12345], [67890]])

    # Should not raise when encountering numeric values
    excel_helper.auto_fit_columns()

    assert excel_helper.active_sheet.column_dimensions["A"].width > 0


def test_apply_style_sets_attributes(excel_helper: ExcelHelper):
    excel_helper.write_cell(1, 1, "Styled")
    excel_helper.apply_style(1, 1, {"font": Font(bold=True)})

    assert excel_helper.read_cell(1, 1) == "Styled"
    assert excel_helper.active_sheet.cell(row=1, column=1).font.bold is True


def test_to_dataframe_empty_sheet_returns_empty_dataframe(excel_helper: ExcelHelper):
    df = excel_helper.to_dataframe()
    assert df.empty


def test_to_dataframe_trims_empty_rows(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [["Header"], ["Value"], [None]])

    df = excel_helper.to_dataframe()

    assert df.shape == (1, 1)
    assert df.iloc[0, 0] == "Value"


def test_to_dataframe_with_sheet_selection(excel_helper: ExcelHelper):
    excel_helper.workbook.create_sheet("Second")
    excel_helper.select_sheet("Second")
    excel_helper.write_range(1, 1, [["Header"], ["Row"]])

    df = excel_helper.to_dataframe(sheet_name="Second")

    assert list(df.columns) == ["Header"]
    assert df.iloc[0, 0] == "Row"


def test_copy_formula_translates_destination(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [[1], [2], [3], [4], [5]])
    excel_helper.sum_range(1, 1, 5, 1, 1, 2)
    excel_helper.copy_formula(1, 2, 2, 2)

    assert excel_helper.get_formula(2, 2) == "=SUM(A2:A6)"


def test_from_dataframe_writes_values(excel_helper: ExcelHelper):
    df = pd.DataFrame({"Col1": [1, 2], "Col2": [3, 4]})
    excel_helper.from_dataframe(df)

    assert excel_helper.read_range(1, 1, 3, 2) == [["Col1", "Col2"], [1, 3], [2, 4]]


def test_add_data_validation_registers_rule(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [["Value"], ["A"]])
    excel_helper.add_data_validation("A2", "list", "between", "\"A,B\"")

    assert len(excel_helper.active_sheet.data_validations.dataValidation) == 1


def test_apply_conditional_formatting_adds_rule(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [["Header"], [1], [2]])
    excel_helper.apply_conditional_formatting("A2:A3", "color_scale")

    assert len(excel_helper.active_sheet.conditional_formatting) == 1


def test_create_chart_adds_bar_chart(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [["Header"], [1], [2]])
    excel_helper.create_chart(
        "bar",
        (1, 1, 1, 3),
        title="Test Chart",
        x_axis="X",
        y_axis="Y",
        location="E1",
    )

    assert len(excel_helper.active_sheet._charts) == 1


def test_freeze_panes_sets_attribute(excel_helper: ExcelHelper):
    excel_helper.freeze_panes("B2")
    assert excel_helper.active_sheet.freeze_panes == "B2"


def test_merge_and_unmerge_cells(excel_helper: ExcelHelper):
    excel_helper.merge_cells(1, 1, 2, 2)
    merged_ranges = list(excel_helper.active_sheet.merged_cells)
    assert len(merged_ranges) == 1

    excel_helper.unmerge_cells(1, 1, 2, 2)
    assert len(list(excel_helper.active_sheet.merged_cells)) == 0


def test_clear_range_removes_values(excel_helper: ExcelHelper):
    excel_helper.write_range(1, 1, [["A", "B"], ["C", "D"]])
    excel_helper.clear_range(1, 1, 2, 2)

    assert excel_helper.read_range(1, 1, 2, 2) == [[None, None], [None, None]]


def test_create_pivot_table_writes_summary(excel_helper: ExcelHelper):
    data = [
        ["Category", "Region", "Sales"],
        ["A", "North", 10],
        ["A", "South", 5],
        ["B", "North", 8],
    ]

    excel_helper.create_pivot_table(
        data,
        pivot_location="E1",
        rows=["Category"],
        columns=[],
        values=["Sales"],
    )

    result = excel_helper.read_range(1, 5, 3, 6)
    assert result[0][:2] == ["Category", "Sales"]
    assert result[1][:2] == ["A", 15]
    assert result[2][:2] == ["B", 8]


def test_use_template_renders_context(excel_helper: ExcelHelper, tmp_path):
    template_path = tmp_path / "template.xlsx"
    output_path = tmp_path / "output.xlsx"

    template_wb = openpyxl.Workbook()
    template_ws = template_wb.active
    template_ws["A1"] = "Hello {{ name }}"
    template_wb.save(template_path)

    excel_helper.use_template(
        str(template_path), str(output_path), {"name": "World"}
    )

    result_wb = openpyxl.load_workbook(output_path)
    assert result_wb.active["A1"].value == "Hello World"
